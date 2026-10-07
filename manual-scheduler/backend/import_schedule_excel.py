"""Import unambiguous timetable cells; dry run by default, never publish.
Usage: python import_schedule_excel.py FILE.xlsx --report REPORT.json [--apply]
Unresolved references and every conflicting lesson stay in the report.
"""
import argparse
from collections import Counter, defaultdict
from contextlib import closing
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import sqlite3
from openpyxl import load_workbook
import main

DAYS = ['понедельник', 'вторник', 'среда', 'четверг', 'пятница', 'суббота']
TYPE_RE = re.compile(r'\s*[-–]\s*(лекция|лек|практ|пр|лаб|студ)\b[\s.,:]*', re.I)
TYPES = {'лекция': 'lecture', 'лек': 'lecture', 'практ': 'practice', 'пр': 'practice', 'лаб': 'lab'}
REASONS = {
 'group': 'Группа не найдена однозначно', 'teacher': 'Преподаватель не найден однозначно',
 'parse': 'Не удалось разделить предмет, тип занятия и преподавателя',
 'studio': 'Студийное занятие: в редакторе нет такого типа',
 'room': 'Кабинет не найден однозначно', 'location': 'Место требует уточнения: онлайн, вне корпуса или не указано',
 'subject': 'Несколько совпадающих предметов', 'conflict': 'Пересечение в исходном Excel',
 'existing_conflict': 'Пересечение с текущим расписанием',
}


def normalize(value):
    return re.sub(r'[^\w]', '', value.casefold().replace('ё', 'е'))


def group_key(value):
    return normalize(re.sub(r'\([^)]*\)', '', value))


def read_excel(path):
    book = load_workbook(path, data_only=True)
    entries = []
    for sheet in book:
        merged = {}
        for area in sheet.merged_cells.ranges:
            for row in range(area.min_row, area.max_row + 1):
                for col in range(area.min_col, area.max_col + 1):
                    merged[row, col] = (area.min_row, area.min_col)
        def value(row, col):
            return str(sheet.cell(*merged.get((row, col), (row, col))).value or '').strip()
        header = next((r for r in range(1, min(20, sheet.max_row) + 1)
                       if any('Дисциплина' in str(c.value) for c in sheet[r])), None)
        if header is None:
            raise ValueError(f'Не найдена строка заголовков на листе {sheet.title}')
        cols = [c.column for c in sheet[header] if 'Дисциплина' in str(c.value)]
        blocks = []
        for i, col in enumerate(cols):
            end = cols[i + 1] if i + 1 < len(cols) else sheet.max_column + 1
            room_col = next((c for c in range(col + 1, end) if 'Ауд' in value(header, c)), None)
            group = value(header - 1, col)
            if group:
                blocks.append((group, col, room_col))
        day = None
        for row in range(header + 1, sheet.max_row + 1):
            for cell in sheet[row]:
                label = str(cell.value or '').strip().lower()
                if label in DAYS:
                    day = DAYS.index(label) + 1
            time = re.fullmatch(r'(\d{1,2})[.:](\d{2})\s*[-–]\s*(\d{1,2})[.:](\d{2})', value(row, 1))
            if time is None or day is None:
                continue
            if time[2] != '00' or time[3] != time[1] or time[4] != '50' or not 8 <= int(time[1]) <= 20:
                raise ValueError(f'Неподдерживаемое время: {sheet.title}!A{row}: {value(row, 1)}')
            for group, col, room_col in blocks:
                description = value(row, col)
                if description:
                    entries.append(dict(sheet=sheet.title, cell=sheet.cell(row, col).coordinate,
                                        group=group, day=day, slot=int(time[1]), text=description,
                                        room=value(row, room_col) if room_col else ''))
    book.close()
    return entries


def resolve(entries, references):
    kinds = {kind: [r for r in references.values() if r['kind'] == kind]
             for kind in ('groups', 'teachers', 'subjects', 'rooms')}
    result = []
    for entry in entries:
        row = dict(entry, issues=[])
        def assign(field, matches, reason):
            if len(matches) == 1:
                row[field] = matches[0]['id']
            else:
                row['issues'].append(reason)
        assign('group_id', [r for r in kinds['groups'] if group_key(r['name']) == group_key(row['group'])], 'group')
        match = TYPE_RE.search(row['text'])
        if not match:
            row['issues'].append('parse')
        else:
            row['subject'] = re.sub(r'\s+', ' ', row['text'][:match.start()]).strip()
            row['lesson_type'] = TYPES.get(match[1].lower())
            if row['lesson_type'] is None:
                row['issues'].append('studio')
            teacher = row['text'][match.end():].strip().split(',')[0]
            teacher = re.split(r'\s+(?:магистр|сеньор|асс|PhD|–)', teacher)[0].strip()
            row['teacher'] = teacher
            hits = []
            for ref in kinds['teachers']:
                tokens = ref['name'].split()
                aliases = {normalize(ref['name'])}
                # Only exact surname + supplied initials; no fuzzy identity guesses.
                for n in range(len(tokens)):
                    aliases.add(normalize(tokens[0] + ''.join(t[0] for t in tokens[1:n + 1])))
                if normalize(teacher) in aliases:
                    hits.append(ref)
            assign('teacher_id', hits, 'teacher')
            subjects = [r for r in kinds['subjects'] if normalize(r['name']) == normalize(row['subject'])]
            exact = [r for r in subjects if r['name'] == row['subject']]
            if len(exact) == 1:
                subjects = exact
            if len(subjects) == 1:
                row['subject_id'] = subjects[0]['id']
            elif len(subjects) > 1:
                row['issues'].append('subject')
        room = re.fullmatch(r'([123])/([\w]+)', row['room'])
        if room:
            assign('room_id', [r for r in kinds['rooms']
                if normalize(r.get('building', '')) == normalize(room[1] + ' блок')
                and normalize(r['name'].split('/')[-1]) == normalize(room[2])], 'room')
        else:
            row['issues'].append('location')
        result.append(row)
    return result


def signature(row):
    subject = ('id', row['subject_id']) if 'subject_id' in row else ('name', normalize(row.get('subject', '')))
    return (row['day'], row['slot'], row.get('teacher_id'), row.get('room_id'), subject, row.get('lesson_type'))


def overlaps(left, right):
    if (left['day'], left['slot']) != (right['day'], right['slot']):
        return False
    return any(left.get(k) is not None and left.get(k) == right.get(k) for k in ('teacher_id', 'room_id')) or bool(set(left.get('group_ids', [])) & set(right.get('group_ids', [])))


def make_plan(entries, references, existing):
    rows = resolve(entries, references)
    candidates = defaultdict(list)
    for row in rows:
        if not row['issues']:
            candidates[signature(row)].append(row)
    accepted = []
    for key, sources in candidates.items():
        first = sources[0]
        item = {k:first[k] for k in ('day', 'slot', 'teacher_id', 'room_id', 'lesson_type', 'subject')}
        if 'subject_id' in first:
            item['subject_id'] = first['subject_id']
        item['group_ids'] = sorted({r['group_id'] for r in sources})
        # Include partly unresolved source rows when checking known occupied
        # teachers, rooms or groups. Do not resolve a conflict by input order.
        conflicts = [r for r in rows if signature(r) != key and overlaps(item, dict(r, group_ids=[r['group_id']] if 'group_id' in r else []))]
        old_conflicts = [r for r in existing if signature(r) != key and overlaps(item, r)]
        if conflicts or old_conflicts:
            for row in sources:
                row['issues'].append('conflict' if conflicts else 'existing_conflict')
                row['conflicts_with'] = [f"{r['sheet']}!{r['cell']}" for r in conflicts]
                row['existing_conflicts'] = [r['id'] for r in old_conflicts]
        else:
            item['sources'] = sources
            accepted.append(item)
    return rows, accepted


def import_file(path, apply=False):
    entries = read_excel(path)
    with main.database(write=apply) as con:
        references, existing = main.refs(con), main.lessons(con)
        rows, plan = make_plan(entries, references, existing)
        backup = None
        created = linked = 0
        if apply:
            folder = main.DB.parent / 'backups'
            folder.mkdir(exist_ok=True)
            backup = folder / ('before-excel-import-' + datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ') + '.sqlite')
            with closing(sqlite3.connect(main.DB.resolve().as_uri() + '?mode=ro', uri=True)) as source, closing(sqlite3.connect(backup)) as target:
                source.backup(target)
            backup.chmod(0o600)
            subject_ids = {}
            previous = {signature(l): l for l in existing}
            for item in plan:
                if 'subject_id' not in item:
                    key = normalize(item['subject'])
                    if key not in subject_ids:
                        data = main.Reference(name=item['subject']).model_dump(exclude={'name', 'expected_version'})
                        subject_ids[key] = con.execute('INSERT INTO refs(kind,name,data) VALUES(?,?,?)',
                            ('subjects', item['subject'], json.dumps(data, ensure_ascii=False))).lastrowid
                    item['subject_id'] = subject_ids[key]
                old = previous.get(signature(item))
                if old:
                    lesson_id = old['id']
                else:
                    lesson_id = con.execute('INSERT INTO lessons(teacher_id,subject_id,room_id,day,slot,lesson_type) VALUES(?,?,?,?,?,?)',
                        tuple(item[k] for k in ('teacher_id','subject_id','room_id','day','slot','lesson_type'))).lastrowid
                    created += 1
                for group_id in item['group_ids']:
                    linked += con.execute('INSERT OR IGNORE INTO lesson_groups(lesson_id,group_id) VALUES(?,?)', (lesson_id,group_id)).rowcount
        summary = dict(source_cells=len(rows), ready_lessons=len(plan), ready_cells=sum(len(p['sources']) for p in plan),
                       skipped_cells=sum(bool(r['issues']) for r in rows), created_lessons=created, added_group_links=linked,
                       sheets=dict(Counter(r['sheet'] for r in rows)), reasons=dict(Counter(REASONS[i] for r in rows for i in r['issues'])))
        return dict(source=Path(path).name, sha256=hashlib.sha256(Path(path).read_bytes()).hexdigest(), applied=apply,
                    backup=str(backup) if backup else None, summary=summary,
                    rows=[dict(r, reasons=[REASONS[i] for i in r['issues']]) for r in rows])


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('--report', required=True, type=Path)
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()
    report = import_file(args.source, args.apply)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2))
    args.report.chmod(0o600)
    print(json.dumps(report['summary'], ensure_ascii=False, indent=2))
    print('Report:', args.report)
