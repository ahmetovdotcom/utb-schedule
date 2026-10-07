import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import Workbook
import main
import import_schedule_excel as importer

class ImportTests(unittest.TestCase):
    def setUp(self):
        self.old_db=main.DB
        self.tmp=tempfile.TemporaryDirectory()
        main.DB=Path(self.tmp.name)/'db.sqlite'
        self.file=Path(self.tmp.name)/'source.xlsx';self.file.write_bytes(b'fixture')
        with main.database(True) as con:
            for kind,name in [('groups','Г-262'),('teachers','Иванов Иван Иванович'),('rooms','1 блок / 302')]:
                value=main.Reference(name=name,course=1,building='1 блок').model_dump(exclude={'name','expected_version'})
                con.execute('insert into refs(kind,name,data) values(?,?,?)',(kind,name,json.dumps(value)))
        self.row=dict(sheet='Кафедра',cell='B14',group='Г-262 (20)',day=1,slot=8,text='Математика-лек. Иванов И.И., профессор',room='1/302')
    def tearDown(self):
        main.DB=self.old_db;self.tmp.cleanup()
    def run_import(self,rows,apply=True):
        with patch.object(importer,'read_excel',return_value=rows):return importer.import_file(self.file,apply)
    def test_dry_run_apply_backup_and_repeat(self):
        self.assertEqual(self.run_import([self.row],False)['summary']['ready_lessons'],1)
        with main.database() as con:self.assertEqual(len(main.lessons(con)),0)
        result=self.run_import([self.row,self.row])
        self.assertTrue(Path(result['backup']).exists())
        self.assertEqual(result['summary']['created_lessons'],1)
        second=self.run_import([self.row])
        self.assertEqual(second['summary']['created_lessons'],0)
        self.assertEqual(second['summary']['added_group_links'],0)
        with main.database() as con:
            self.assertEqual(len(main.lessons(con)),1)
            self.assertEqual(con.execute("select count(*) from refs where kind='subjects'").fetchone()[0],1)
    def test_conflicts_and_unknowns_are_reported_without_writes(self):
        different=dict(self.row,text='Физика-лек. Иванов И.И.',cell='F14')
        report=self.run_import([self.row,different,dict(self.row,group='Неизвестная',slot=9)])
        self.assertEqual(report['summary']['created_lessons'],0)
        self.assertEqual(report['summary']['skipped_cells'],3)
        self.run_import([self.row])
        report=self.run_import([different])
        self.assertIn('existing_conflict',report['rows'][0]['issues'])
        self.assertEqual(report['summary']['created_lessons'],0)
    def test_remote_import_is_idempotent(self):
        row = dict(self.row, room='902 876 6192, код: DD57YJ')
        self.assertEqual(self.run_import([row])['summary']['created_lessons'],1)
        self.assertEqual(self.run_import([row])['summary']['created_lessons'],0)
        with main.database() as con:
            lesson=main.lessons(con)[0]
            self.assertEqual(lesson['delivery'],'online')
            self.assertIsNone(lesson['room_id'])
            self.assertEqual(lesson['meeting_id'],'9028766192')
            self.assertEqual(lesson['passcode'],'DD57YJ')

    def test_reader_handles_short_blocks_and_merged_rooms(self):
        book=Workbook();sheet=book.active
        sheet['B10']='Г-262';sheet['D10']='Г-263'
        sheet['B11']='Дисциплина';sheet['C11']='Ауд.';sheet['D11']='Дисциплина';sheet['G11']='Ауд.'
        sheet['B12']='ПОНЕДЕЛЬНИК'
        for row in (13,14):
            sheet.cell(row,1,f'{row-5:02d}.00-{row-5:02d}.50')
            sheet.cell(row,2,self.row['text'])
        sheet['C13']='1/302';sheet.merge_cells('C13:C14')
        with patch.object(importer,'load_workbook',return_value=book):rows=importer.read_excel(self.file)
        self.assertEqual(len(rows),2)
        self.assertEqual([r['room'] for r in rows],['1/302','1/302'])
        self.assertEqual([r['slot'] for r in rows],[8,9])

if __name__=='__main__':unittest.main()
