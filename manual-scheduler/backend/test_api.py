import tempfile
import unittest
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import main
import auth

class ScheduleTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        main.DB = Path(self.temp.name) / 'test.db'
        auth.configure_account('admin','test-password-12345')
        self.client = TestClient(main.app)
        self.client.post('/api/auth/login',json={'login':'admin','password':'test-password-12345'})
        self.assertEqual(self.client.post('/api/demo').status_code, 200)
        state = self.client.get('/api/state').json()
        self.refs = state['references']
        self.lesson = state['lessons'][0]
        self.payload = {k:v for k,v in self.lesson.items() if k!='id'}
        self.payload['slot'] = 10
        self.payload['expected_version'] = self.lesson['version']

    def tearDown(self):
        self.temp.cleanup()

    def test_conflicts_and_stream_warnings(self):
        groups = [r['id'] for r in self.refs if r['kind']=='groups']
        self.payload['group_ids'] = groups
        response = self.client.post('/api/lessons',json=self.payload)
        self.assertEqual(response.status_code,409)
        self.assertTrue(response.json()['detail']['warnings'])
        response = self.client.post('/api/lessons',json={**self.payload,'accept_warnings':True})
        self.assertEqual(response.status_code,201)
        response = self.client.post('/api/lessons',json={**self.payload,'accept_warnings':True})
        self.assertEqual(response.status_code,409)
        self.assertTrue(response.json()['detail']['conflicts'])
        self.assertEqual(len(self.client.get('/api/state').json()['lessons']),2)

    def test_update_excludes_itself_and_delete_cascades(self):
        lid = self.lesson['id']
        response=self.client.put(f'/api/lessons/{lid}',json=self.payload)
        self.assertEqual(response.status_code,200)
        options=self.client.post(f'/api/options?exclude={lid}',json=self.payload).json()
        current=next(o for o in options if o['room_id']==self.payload['room_id'] and o['slot']==10)
        self.assertFalse(current['conflicts'])
        self.assertEqual(self.client.delete(f'/api/lessons/{lid}',params={'expected_version':self.client.get('/api/state').json()['lessons'][0]['version']}).status_code,200)
        self.assertFalse(self.client.get('/api/state').json()['lessons'])

    def test_invalid_references_and_referenced_delete(self):
        self.assertEqual(self.client.post('/api/lessons',json={**self.payload,'teacher_id':self.payload['subject_id']}).status_code,422)
        self.assertEqual(self.client.delete('/api/references/'+str(self.lesson['room_id']),params={'expected_version':next(r['version'] for r in self.refs if r['id']==self.lesson['room_id'])}).status_code,409)
        self.assertEqual(self.client.post('/api/demo').status_code,409)

    def test_concurrent_booking(self):
        def create(_):
            with TestClient(main.app) as client:
                client.cookies.update(self.client.cookies)
                return client.post('/api/lessons',json=self.payload).status_code
        with ThreadPoolExecutor(max_workers=2) as executor:
            self.assertEqual(sorted(executor.map(create,range(2))),[201,409])

    def test_conflicts_across_rooms_and_reference_edit(self):
        other_room = next(r for r in self.refs if r['kind']=='rooms' and r['name']=='202')
        response = self.client.post('/api/lessons',json={**self.payload,'slot':9,'room_id':other_room['id'],'accept_warnings':True})
        self.assertEqual(response.status_code,409)
        self.assertIn('Преподаватель занят',response.json()['detail']['conflicts'][0]['message'])
        self.assertIn('Группа занята',response.json()['detail']['conflicts'][0]['message'])
        group = next(r for r in self.refs if r['id']==self.payload['group_ids'][0])
        response = self.client.put('/api/references/'+str(group['id']),json={**group,'students':200,'expected_version':group['version']})
        self.assertEqual(response.status_code,200)
        self.assertTrue(self.client.get('/api/state').json()['lessons'][0]['warnings'])

    def test_export_course_and_literal_cells(self):
        with TestClient(main.app) as anonymous:
            self.assertEqual(anonymous.get('/api/export.xlsx?course=2').status_code,401)
        self.client.post('/api/references/groups',json={'name':'=1+1','course':2,'department':'Другая'})
        response=self.client.get('/api/export.xlsx?course=2')
        self.assertEqual(response.status_code,200)
        book=load_workbook(BytesIO(response.content))
        self.assertEqual(set(book.sheetnames),{'Химия','Другая'})
        self.assertEqual(book['Другая']['B11'].value,'=1+1')
        self.assertEqual(book['Другая']['B11'].data_type,'s')
        self.assertIn('Химия',book['Химия']['B15'].value)
        self.assertEqual(self.client.get('/api/export.xlsx?course=8').status_code,404)

if __name__=='__main__': unittest.main()
