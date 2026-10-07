from io import BytesIO
from copy import copy
import unittest
from openpyxl import load_workbook
from excel_export import build_workbook, TEMPLATE


class ExcelTests(unittest.TestCase):
    def test_template_streams_empty_groups_and_extended_times(self):
        refs = {
            1: dict(id=1, kind='subjects', name='=Предмет'),
            2: dict(id=2, kind='teachers', name='Преподаватель'),
            3: dict(id=3, kind='rooms', name='406', building='1'),
        }
        for index in range(10, 20):
            refs[index] = dict(id=index, kind='groups', name=f'Г-{index}', students=20, course=3, department='Кафедра')
        refs[20] = dict(id=20, kind='groups', name='Без кафедры группа', students=10, course=3, department='  ')
        refs[21] = dict(id=21, kind='groups', name='Другой курс', students=10, course=2, department='Кафедра')
        lessons = [dict(group_ids=[10,19,20,21], day=6, slot=20, subject_id=1, teacher_id=2, room_id=3, lesson_type='lecture')]
        book = load_workbook(BytesIO(build_workbook(3, refs, lessons)))
        self.assertEqual(book.sheetnames, ['Кафедра', 'Без кафедры'])
        original = load_workbook(TEMPLATE).worksheets[0]
        for sheet in book:
            self.assertEqual(sheet['G1'].value, original['G1'].value)
            self.assertEqual(copy(sheet['G1'].font), copy(original['G1'].font))
            self.assertEqual(sheet.page_setup.orientation, 'landscape')
            self.assertEqual(sheet['B12'].value, original['B12'].value)
            self.assertIn('СУББОТА', [c.value for c in sheet['A']])
            self.assertFalse(any('1 С Предприятие' in str(c.value) or 'АЖ-241' in str(c.value) or 'Другой курс' in str(c.value) for row in sheet for c in row))
            target = next(c.row for c in sheet['A'] if c.value == '20.00-20.50')
            self.assertEqual(sheet.cell(target,2).value, '=Предмет-лек. Преподаватель')
            self.assertEqual(sheet.cell(target,2).data_type, 's')
            self.assertEqual(sheet.cell(target,5).value, '1/406')
        sheet = book['Кафедра']
        self.assertEqual(sheet['AL11'].value, 'Г-19 (20)')
        target = next(c.row for c in sheet['A'] if c.value == '20.00-20.50')
        self.assertEqual(sheet.cell(target,38).value, '=Предмет-лек. Преподаватель')
        self.assertIsNone(sheet.cell(target,6).value)
        self.assertIn('AO', str(sheet.print_area))

    def test_sheet_name_collisions_and_unknown_course(self):
        names = ['a/b', 'a?b', 'History', 'Без кафедры', '', 'X'*40, 'x'*40]
        refs = {i:dict(id=i, kind='groups', name=str(i), course=0, students=20, department=name) for i,name in enumerate(names)}
        book = load_workbook(BytesIO(build_workbook(0,refs,[])))
        self.assertEqual(len(book.sheetnames), len(names))
        self.assertEqual(len({s.casefold() for s in book.sheetnames}), len(names))
        self.assertTrue(all(len(s)<=31 for s in book.sheetnames))
        self.assertEqual(book.worksheets[-1]['F4'].value.split(' на базе')[0], 'Курс не указан')
        with self.assertRaisesRegex(ValueError, 'нет групп'):
            build_workbook(8,refs,[])
