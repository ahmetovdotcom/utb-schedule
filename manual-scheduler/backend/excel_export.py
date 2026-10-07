"""Render the editor snapshot using the supplied university Excel template."""
from collections import defaultdict
from copy import copy, deepcopy
from io import BytesIO
from pathlib import Path
import re
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

TEMPLATE = Path(__file__).with_name('templates') / 'timetable.xlsx'
DAYS = ['ПОНЕДЕЛЬНИК', 'ВТОРНИК', 'СРЕДА', 'ЧЕТВЕРГ', 'ПЯТНИЦА', 'СУББОТА']
TYPES = {'lecture': 'лек.', 'practice': 'пр.', 'lab': 'лаб.'}


def text(sheet, row, col, value):
    cell = sheet.cell(row, col)
    cell.value = ILLEGAL_CHARACTERS_RE.sub('', str(value))
    cell.data_type = 's'


def sheet_name(name, used):
    base = re.sub(r'[\\/*?:\[\]]', ' ', ILLEGAL_CHARACTERS_RE.sub('', name)).strip().strip("'")[:31] or 'Кафедра'
    if base.casefold() == 'history':
        base = 'History кафедра'
    result, index = base, 2
    while result.casefold() in used:
        suffix = f' ({index})'
        result = base[:31-len(suffix)] + suffix
        index += 1
    used.add(result.casefold())
    return result


def build_workbook(course, references, lessons):
    book = load_workbook(TEMPLATE)
    source = book.worksheets[0]
    departments = defaultdict(list)
    for group in references.values():
        if group['kind'] == 'groups' and group.get('course') == course:
            departments[group.get('department', '').strip()].append(group)
    if not departments:
        raise ValueError('Для этого курса нет групп')
    headings = {DAYS.index(str(source.cell(r, 1).value)) + 1: r
                for r in range(13, source.max_row + 1) if str(source.cell(r, 1).value) in DAYS}
    footer = max(r for r in range(13, source.max_row + 1)
                 if re.fullmatch(r'\d{2}\.\d{2}-\d{2}\.\d{2}', str(source.cell(r, 1).value))) + 1
    by_cell = defaultdict(list)
    for lesson in lessons:
        for group_id in lesson['group_ids']:
            by_cell[group_id, lesson['day'], lesson['slot']].append(lesson)
    # Remove the example sheets only after creating output sheets.
    originals, used, output_sheets = list(book.worksheets), set(), []
    for department, groups in sorted(departments.items(), key=lambda item: (not item[0], item[0].casefold())):
        groups.sort(key=lambda group: (group['name'].casefold(), group['id']))
        title = sheet_name(department or 'Без кафедры', used)
        sheet = book.create_sheet(uuid4().hex[:31])
        output_sheets.append((sheet, title))
        last_col = max(33, 1 + 4 * len(groups))
        if last_col > 16384:
            raise ValueError('Слишком много групп для одного листа Excel')
        ids = {g['id'] for g in groups}
        relevant = [l for l in lessons if ids.intersection(l['group_ids'])]
        plan = [(r, None, None) for r in range(1, 13)]
        for day in range(1, 7 if any(l['day'] == 6 for l in relevant) else 6):
            heading = headings.get(day, headings[1])
            end = headings.get(day + 1, footer) if day in headings else headings[2]
            time_rows = {int(str(source.cell(r, 1).value)[:2]): r for r in range(heading + 1, end)
                         if re.fullmatch(r'\d{2}\.\d{2}-\d{2}\.\d{2}', str(source.cell(r, 1).value))}
            plan.append((heading, day, None))
            required = {l['slot'] for l in relevant if l['day'] == day}
            for hour in sorted(set(time_rows) | required):
                plan.append((time_rows.get(hour, 14), day, hour))
            if day in headings:
                plan.extend((r, day, -1) for r in range(heading + 1, end) if not source.cell(r, 1).value)
        plan.extend((r, None, None) for r in range(footer, source.max_row + 1))
        origin = {}
        for row, (src_row, day, hour) in enumerate(plan, 1):
            if day is None:
                origin[src_row] = row
            sheet.row_dimensions[row] = copy(source.row_dimensions[src_row])
            sheet.row_dimensions[row].index = row
            for col in range(1, max(last_col, source.max_column) + 1):
                src_col = col if col <= 33 else 2 + (col - 34) % 4
                cell = sheet.cell(row, col)
                cell._style = copy(source.cell(src_row, src_col)._style)
                if day is None and (src_row < 11 or src_row >= footer) and col <= source.max_column:
                    value = source.cell(src_row, col).value
                    if value is not None:
                        text(sheet, row, col, value)
            if day is not None:
                if hour is None:
                    text(sheet, row, 1, DAYS[day - 1])
                    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
                else:
                    if hour >= 0:
                        text(sheet, row, 1, f'{hour:02d}.00-{hour:02d}.50')
                    for col in range(2, last_col + 1, 4):
                        sheet.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 2)
                    for index, group in enumerate(groups):
                        entries = by_cell[group['id'], day, hour]
                        descriptions, rooms = [], []
                        for lesson in entries:
                            descriptions.append(f"{references[lesson['subject_id']]['name']}-{TYPES[lesson['lesson_type']]} {references[lesson['teacher_id']]['name']}")
                            room = references[lesson['room_id']]
                            rooms.append('/'.join(filter(None, [room.get('building', ''), room['name']])))
                        text(sheet, row, 2 + 4 * index, '\n'.join(descriptions))
                        text(sheet, row, 5 + 4 * index, '\n'.join(rooms))
        for merged in source.merged_cells.ranges:
            if merged.min_row in origin and merged.max_row in origin and merged.min_row != 11 and merged.min_row != 12:
                right = last_col - 1 if merged.max_col == 32 and merged.min_row <= 5 else merged.max_col
                sheet.merge_cells(start_row=origin[merged.min_row], start_column=merged.min_col,
                                  end_row=origin[merged.max_row], end_column=right)
        sheet.merge_cells('A11:A12')
        text(sheet, 11, 1, source['A11'].value)
        for index, col in enumerate(range(2, last_col + 1, 4)):
            sheet.merge_cells(start_row=11, start_column=col, end_row=11, end_column=col + 3)
            sheet.merge_cells(start_row=12, start_column=col, end_row=12, end_column=col + 2)
            text(sheet, 12, col, source['B12'].value)
            text(sheet, 12, col + 3, source['E12'].value)
            if index < len(groups):
                group = groups[index]
                text(sheet, 11, col, f"{group['name']} ({group['students']})")
        text(sheet, 4, 6, re.sub(r'^\d+ курс', f'{course} курс' if course else 'Курс не указан', source['F4'].value))
        text(sheet, 5, 6, re.sub(r'^\d+ семестр', f'{course * 2 - 1} семестр' if course else 'Семестр не указан', source['F5'].value))
        for key, dimension in source.column_dimensions.items():
            sheet.column_dimensions[key] = copy(dimension)
        for col in range(34, last_col + 1):
            sheet.column_dimensions[get_column_letter(col)].width = source.column_dimensions['B'].width
        for attr in ('sheet_format', 'sheet_properties', 'views', 'page_setup', 'page_margins', 'print_options', 'HeaderFooter'):
            setattr(sheet, attr, deepcopy(getattr(source, attr)))
        sheet.print_area = f'A1:{get_column_letter(last_col)}{len(plan)}'
        sheet.print_title_rows = '11:12'
        sheet.freeze_panes = 'B13'
    for sheet in originals:
        book.remove(sheet)
    for sheet, title in output_sheets:
        sheet.title = title
    book.active = 0
    output = BytesIO()
    book.save(output)
    return output.getvalue()
